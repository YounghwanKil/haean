import io
import json
import zipfile

import openpyxl

from haean.corpus import Corpus


def test_xlsx_merged_cells_formula_and_dedup(tmp_path):
    # Synthetic source fixture only; this does not author a user workbook.
    path = tmp_path / "PSAT_7급_test.xlsx"
    w = openpyxl.Workbook()
    s = w.active
    s.title = "문항 데이터"
    s.append(["통합 문항 ID", "유형", "비고"])
    s.append(["2026-VL-01", "독해", "=1+1"])
    other = w.create_sheet("merged")
    other.merge_cells("A1:C1")
    other["A1"] = "header"
    w.save(path)
    c = Corpus(tmp_path / "corpus.sqlite")
    assert c.import_file(path)["status"] == "ok"
    assert c.import_file(path)["status"] == "duplicate"
    rows = c.rows("psat7")
    assert next(r for r in rows if r["row"] == 2)["uncached_formulas"] == ["C2"]
    assert c.db.execute("select count(*) from records where role='metadata'").fetchone()[0] == 1


def test_archive_cannot_write_or_execute(tmp_path):
    path = tmp_path / "archive.zip"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("../../escape.md", "reference only")
        z.writestr("run.sh", "touch /tmp/should-never-run")
    c = Corpus(tmp_path / "db")
    assert c.import_file(path)["status"] == "ok"
    assert not (tmp_path.parent / "escape.md").exists()
    assert c.db.execute("select count(*) from records").fetchone()[0] == 0


def test_corrupt_import_rolls_back(tmp_path):
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"not a workbook")
    c = Corpus(tmp_path / "db")
    assert c.import_file(path)["status"] == "error"
    assert c.db.execute("select count(*) from records").fetchone()[0] == 0


def test_renamed_identical_file_keeps_source_id(tmp_path):
    one = Corpus(tmp_path / "a.db")
    two = Corpus(tmp_path / "b.db")
    one.add("samehash", "name.xlsx#문항 데이터!2", "leet", "example", "text")
    two.add("samehash", "renamed.xlsx#문항 데이터!2", "leet", "example", "text")
    assert one.db.execute("select id from records").fetchone()[0] == two.db.execute("select id from records").fetchone()[0]

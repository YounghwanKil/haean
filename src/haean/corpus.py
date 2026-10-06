"""Read-only source import. Archives are read recursively; embedded instructions are never executed."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import sqlite3
import struct
import unicodedata
import xml.etree.ElementTree as ET
import zipfile
import zlib
from pathlib import Path

import olefile
import openpyxl


def nfc(value: str) -> str:
    return unicodedata.normalize("NFC", value)


def decode_text(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp949")


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def archive_name(info: zipfile.ZipInfo) -> str:
    name = info.filename
    if not info.flag_bits & 0x800:
        try:
            name = name.encode("cp437").decode("utf-8")
        except UnicodeError:
            pass
    return nfc(name)


def hwp_text(data: bytes) -> str:
    """Extract HWP5 paragraph records; formulas/images/layout remain unverified."""
    with olefile.OleFileIO(io.BytesIO(data)) as doc:
        header = doc.openstream("FileHeader").read()
        flags = struct.unpack_from("<I", header, 36)[0]
        if flags & (2 | 4):
            raise ValueError("Encrypted/distribution HWP needs a supported converter")
        paragraphs = []
        sections = sorted((p for p in doc.listdir() if p[0] == "BodyText"),
                          key=lambda p: int(re.search(r"\d+$", p[-1]).group()))
        for section in sections:
            raw = doc.openstream(section).read()
            if flags & 1:
                raw = zlib.decompress(raw, -15)
            offset = 0
            while offset + 4 <= len(raw):
                word = struct.unpack_from("<I", raw, offset)[0]
                offset += 4
                tag, size = word & 0x3FF, word >> 20
                if size == 0xFFF:
                    size = struct.unpack_from("<I", raw, offset)[0]
                    offset += 4
                if offset + size > len(raw):
                    raise ValueError("Truncated HWP record")
                if tag == 67:
                    payload = raw[offset:offset + size]
                    units = struct.unpack("<" + "H" * (len(payload) // 2), payload)
                    chars, i = [], 0
                    extended = {1, 2, 3, 4, 5, 6, 7, 8, 9, 11, 12, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23}
                    while i < len(units):
                        code = units[i]
                        if code in extended:
                            chars.append(" " if code == 9 else " [개체] ")
                            i += 8
                        else:
                            chars.append(chr(code) if code >= 32 else "\n" if code in {10, 13} else "")
                            i += 1
                    paragraphs.append("".join(chars))
                offset += size
        if not paragraphs:
            raise ValueError("No HWP paragraph records extracted")
        return "\n".join(paragraphs)


def xml_text(data: bytes) -> str:
    root = ET.fromstring(data)
    return "\n".join("".join(e.itertext()) for e in root.iter()
                     if e.tag.rsplit("}", 1)[-1] in {"t", "script"})


class Corpus:
    def __init__(self, path: Path | str = "data/corpus.sqlite"):
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.db = sqlite3.connect(path)
        self.db.row_factory = sqlite3.Row
        self.db.executescript("""
        CREATE TABLE IF NOT EXISTS sources(
          sha TEXT PRIMARY KEY, name TEXT NOT NULL, path TEXT NOT NULL,
          status TEXT NOT NULL, detail TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS records(
          id TEXT PRIMARY KEY, sha TEXT NOT NULL, locator TEXT NOT NULL,
          exam TEXT NOT NULL, role TEXT NOT NULL, text TEXT NOT NULL, meta TEXT NOT NULL);
        CREATE INDEX IF NOT EXISTS record_exam ON records(exam,role);
        """)

    def add(self, sha, locator, exam, role, text, meta=None):
        stable_locator = locator.split("::", 1)[1] if "::" in locator else locator.split("#", 1)[-1]
        rid = digest((sha + "#" + stable_locator).encode())[:24]
        self.db.execute("INSERT OR REPLACE INTO records VALUES (?,?,?,?,?,?,?)",
                        (rid, sha, locator, exam, role, nfc(text), json.dumps(meta or {}, ensure_ascii=False, default=str)))

    def import_file(self, path: Path, refresh=False) -> dict:
        name = nfc(path.name)
        if "폰트모음" in name:
            with path.open("rb") as stream:
                sha = hashlib.file_digest(stream, "sha256").hexdigest()
            data = b""
        else:
            data = path.read_bytes()
            sha = digest(data)
        prior = self.db.execute("SELECT status FROM sources WHERE sha=?", (sha,)).fetchone()
        if prior and prior[0] != "error" and not refresh:
            return {"name": name, "status": "duplicate", "sha": sha}
        exam = "psat7" if "PSAT_7" in name else "psat5" if "PSAT_5" in name else "leet"
        status, detail = "ok", ""
        self.db.execute("SAVEPOINT importing")
        try:
            if refresh:
                self.db.execute("DELETE FROM records WHERE sha=?", (sha,))
            if "폰트모음" in name:
                from .source_audit import archive_inventory
                members = archive_inventory(path)
                status, detail = "asset_only", json.dumps({"members": len(members), "font_payloads": sum(m["status"] == "inventoried" for m in members), "installed": False})
            else:
                self._extract(data, sha, name, exam)
                if path.suffix.lower() in {".hwp", ".hwpx"}:
                    status, detail = "text_only", "Tables, equations, figures and layout need visual verification"
            self.db.execute("RELEASE importing")
        except Exception as exc:
            self.db.execute("ROLLBACK TO importing")
            self.db.execute("RELEASE importing")
            status, detail = "error", f"{type(exc).__name__}: {exc}"
        self.db.execute("INSERT OR REPLACE INTO sources VALUES (?,?,?,?,?)",
                        (sha, name, str(path.resolve()), status, detail))
        self.db.commit()
        return {"name": name, "status": status, "detail": detail, "sha": sha}

    def _extract(self, data, sha, name, exam):
        suffix = Path(name).suffix.lower()
        if suffix == ".xlsx":
            values = openpyxl.load_workbook(io.BytesIO(data), read_only=False, data_only=True)
            formulas = openpyxl.load_workbook(io.BytesIO(data), read_only=False, data_only=False)
            for sheet in values:
                headers = [str(c.value or openpyxl.utils.get_column_letter(i)) for i, c in enumerate(sheet[1], 1)]
                for row in sheet:
                    cells, fields, uncached = {}, {}, []
                    for cell in row:
                        formula = formulas[sheet.title][cell.coordinate]
                        if cell.value is not None or formula.data_type == "f" or formula.comment:
                            cells[cell.coordinate] = {"value": cell.value}
                            if formula.data_type == "f":
                                cells[cell.coordinate]["formula"] = formula.value
                                if cell.value is None:
                                    uncached.append(cell.coordinate)
                            if formula.comment:
                                cells[cell.coordinate]["comment"] = formula.comment.text
                            fields[headers[cell.column - 1]] = cell.value
                    if not cells:
                        continue
                    role = "metadata" if sheet.title == "문항 데이터" and row[0].row > 1 and exam.startswith("psat") else "reference"
                    if sheet.title == "문항 데이터" and row[0].row > 1 and exam == "leet":
                        role = "example" if fields.get("지문") else "metadata"
                    if "의견" in name or "검토의견" in sheet.title:
                        role = "feedback"
                    self.add(sha, f"{name}#{sheet.title}!{row[0].row}", exam, role,
                             json.dumps({"sheet": sheet.title, "fields": fields, "cells": cells}, ensure_ascii=False, default=str),
                             {"sheet": sheet.title, "row": row[0].row, "fields": fields,
                              "uncached_formulas": uncached, "label_status": "source_unverified"})
            values.close()
            formulas.close()
        elif suffix == ".zip":
            member_report = []
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                for info in archive.infolist():
                    member = archive_name(info)
                    parts = Path(member).parts
                    result = {"member": member, "bytes": info.file_size}
                    member_report.append(result)
                    if info.is_dir() or "__MACOSX" in parts or any(p.startswith(".") for p in parts):
                        result["status"] = "metadata_or_runtime_excluded"
                        continue
                    if Path(member).is_absolute() or ".." in parts or "\\" in member:
                        result["status"] = "unsafe_path"
                        continue
                    if info.file_size > 20_000_000:
                        result["status"] = "oversized_member"
                        continue
                    if "전체(단일파일)" in member:
                        result["status"] = "compiled_duplicate"
                        continue
                    nested = name + "::" + member
                    ext = Path(member).suffix.lower()
                    if ext == ".md":
                        text = archive.read(info).decode("utf-8-sig")
                        role = "feedback" if "대조정리" in name else "wiki"
                        self._chunks(sha, nested, exam, role, text)
                    elif ext in {".txt", ".json", ".xlsx", ".hwp", ".hwpx"}:
                        self._extract(archive.read(info), sha, nested, exam)
                    else:
                        result["status"] = "unsupported_format"
                        continue
                    result["status"] = "text_indexed_visual_unverified" if ext in {".hwp", ".hwpx"} else "indexed"
            self.db.execute("CREATE TABLE IF NOT EXISTS archive_members(sha TEXT, member TEXT, detail TEXT, PRIMARY KEY(sha,member))")
            self.db.execute("DELETE FROM archive_members WHERE sha=?", (sha,))
            for result in member_report:
                self.db.execute("INSERT OR REPLACE INTO archive_members VALUES (?,?,?)",
                                (sha, result["member"], json.dumps(result, ensure_ascii=False)))
        elif suffix == ".hwp":
            self._chunks(sha, name, exam, "feedback" if "검토의견" in name else "document", hwp_text(data), {"visual_verified": False})
        elif suffix in {".hwpx", ".docx"}:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                members = sorted(n for n in z.namelist() if re.match(r"Contents/section\d+\.xml$", n) or n == "word/document.xml")
                for member in members:
                    self._chunks(sha, name + "::" + member, exam, "document", xml_text(z.read(member)), {"visual_verified": False})
        elif suffix == ".csv":
            text = decode_text(data)
            for i, row in enumerate(csv.reader(io.StringIO(text)), 1):
                self.add(sha, f"{name}#row{i}", "team", "conversation", " | ".join(row))
        elif suffix == ".json":
            value = json.loads(decode_text(data))
            self._chunks(sha, name, exam, "reference", json.dumps(value, ensure_ascii=False, indent=2))
        elif suffix in {".md", ".txt"}:
            self._chunks(sha, name, exam, "reference", decode_text(data))
        else:
            raise ValueError(f"Unsupported source: {suffix}")

    def _chunks(self, sha, name, exam, role, text, meta=None):
        for i, start in enumerate(range(0, len(text), 5500)):
            self.add(sha, f"{name}#chunk{i + 1}", exam, role, text[start:start + 6000], meta)

    def rows(self, exam, sheet="문항 데이터"):
        result = []
        for row in self.db.execute("SELECT * FROM records WHERE exam=?", (exam,)):
            meta = json.loads(row["meta"])
            if meta.get("sheet") == sheet:
                result.append({"source_id": row["id"], "locator": row["locator"], **meta})
        return result

    def search(self, query, exam=None, limit=8, roles=None):
        terms = re.findall(r"[\w가-힣]+", nfc(query).lower())
        found = []
        for row in self.db.execute("SELECT * FROM records"):
            if exam and row["exam"] != exam:
                continue
            if roles and row["role"] not in roles:
                continue
            text = row["text"].lower()
            score = sum(min(text.count(t), 5) for t in terms)
            if score:
                found.append((score, dict(row)))
        found.sort(key=lambda x: (-x[0], x[1]["locator"]))
        return [dict(row, score=score) for score, row in found[:limit]]

    def report(self):
        return {"sources": [dict(r) for r in self.db.execute("SELECT * FROM sources")],
                "counts": [dict(r) for r in self.db.execute("SELECT exam,role,count(*) AS count FROM records GROUP BY exam,role")]}
